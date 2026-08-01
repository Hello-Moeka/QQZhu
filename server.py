import aiohttp_jinja2
import jinja2
from aiohttp import web
import configparser
import asyncio
import aiohttp
import aiosqlite
import json
import os
import secrets
import time
import io
import math
import zipfile
from PIL import Image, ImageDraw, ImageFont
from urllib.parse import quote, urlparse
import ipaddress
from pypinyin import pinyin, Style
from dotenv import load_dotenv

load_dotenv()

# ---------------------------------------------------------------------------
# 默认设置
# ---------------------------------------------------------------------------
DEFAULT_SETTINGS = {
    "title": "歌单",
    "live_url": "https://",
    "background_url": "/static/background.jpg",
    "singer_url": "/static/singer.jpg",
    "singer_name": "歌手名字",
    "singer_intro": "这里填写歌手介绍~",
}

# ---------------------------------------------------------------------------
# 配置
# ---------------------------------------------------------------------------
class CaseSensitiveConfigParser(configparser.ConfigParser):
    def optionxform(self, optionstr):
        return optionstr


class Config:
    """负责读取和管理配置文件的类。"""

    def __init__(self, filename="config.ini"):
        self.config = CaseSensitiveConfigParser()
        self.config.read(filename)
        self.env_admin_token = os.environ.get("QQZHU_ADMIN_TOKEN")
        self.env_super_admin_token = os.environ.get("QQZHU_SUPER_ADMIN_TOKEN")
        self.single_tenant = os.environ.get("QQZHU_SINGLE_TENANT", "").lower() in ("1", "true", "yes")

    def _get(self, section, key, default):
        if self.config.has_section(section) and key in self.config[section]:
            return self.config[section][key]
        return default

    def db_path(self):
        return os.environ.get("QQZHU_DB_PATH") or self._get("database", "path", "instance/qqzhu.db")

    def server_port(self):
        env_port = os.environ.get("QQZHU_PORT")
        if env_port:
            return int(env_port)
        return int(self._get("server", "port", 8080))

    def admin_token(self):
        return self.env_admin_token or self._get("server", "admin_token", "")

    def super_admin_token(self):
        return self.env_super_admin_token or self._get("super_admin", "token", "")


# ---------------------------------------------------------------------------
# 登录限流
# ---------------------------------------------------------------------------
LOGIN_ATTEMPTS = {}
LOGIN_LIMIT = 5
LOGIN_WINDOW = 300  # seconds
LOGIN_GC_INTERVAL = 600  # 清理周期（秒）
_LOGIN_LAST_GC = 0.0


def _gc_login_attempts():
    """定期清理过期 IP 条目，避免 LOGIN_ATTEMPTS 无限增长。"""
    global _LOGIN_LAST_GC
    now = time.time()
    if now - _LOGIN_LAST_GC < LOGIN_GC_INTERVAL:
        return
    _LOGIN_LAST_GC = now
    expired = [ip for ip, ts in LOGIN_ATTEMPTS.items() if not ts or now - ts[-1] >= LOGIN_WINDOW]
    for ip in expired:
        LOGIN_ATTEMPTS.pop(ip, None)


def _login_attempts_info(ip: str):
    _gc_login_attempts()
    now = time.time()
    attempts = LOGIN_ATTEMPTS.get(ip, [])
    attempts = [t for t in attempts if now - t < LOGIN_WINDOW]
    LOGIN_ATTEMPTS[ip] = attempts
    remaining = max(LOGIN_LIMIT - len(attempts), 0)
    wait = 0
    if remaining == 0 and attempts:
        wait = max(0, int(LOGIN_WINDOW - (now - attempts[0])))
    return remaining, wait


def _add_login_failure(ip: str):
    _gc_login_attempts()
    now = time.time()
    attempts = LOGIN_ATTEMPTS.get(ip, [])
    attempts = [t for t in attempts if now - t < LOGIN_WINDOW]
    attempts.append(now)
    LOGIN_ATTEMPTS[ip] = attempts


def _reset_login_failure(ip: str):
    LOGIN_ATTEMPTS.pop(ip, None)


def update_env_var(key: str, value: str, env_path: str = ".env"):
    """Persist a key-value to .env, replacing if exists."""
    lines = []
    found = False
    if os.path.exists(env_path):
        with open(env_path, "r", encoding="utf-8") as f:
            for line in f:
                if line.strip().startswith(f"{key}="):
                    lines.append(f"{key}={value}\n")
                    found = True
                else:
                    lines.append(line)
    if not found:
        lines.append(f"{key}={value}\n")
    with open(env_path, "w", encoding="utf-8") as f:
        f.writelines(lines)
    os.environ[key] = value


# ---------------------------------------------------------------------------
# 中间件
# ---------------------------------------------------------------------------
@web.middleware
async def csrf_middleware(request, handler):
    is_admin_path = request.path.startswith("/admin") or request.path.startswith("/super-admin")
    if is_admin_path:
        token = request.cookies.get("csrf_token") or secrets.token_urlsafe(32)
        request["csrf_token"] = token

        if request.method == "POST":
            try:
                form = await request.post()
            except Exception:
                form = {}
            form_token = form.get("csrf_token") if hasattr(form, "get") else None
            header_token = request.headers.get("X-CSRF-Token")
            cookie_token = request.cookies.get("csrf_token")
            final_token = form_token or header_token
            if not cookie_token or not final_token or final_token != cookie_token:
                raise web.HTTPForbidden(text="Invalid CSRF token")

    response = await handler(request)

    if is_admin_path and "csrf_token" in request:
        if request.cookies.get("csrf_token") != request["csrf_token"]:
            response.set_cookie("csrf_token", request["csrf_token"], httponly=True, samesite="Lax")
    return response


@web.middleware
async def security_headers_middleware(request, handler):
    response = await handler(request)
    response.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault(
        "Content-Security-Policy",
        "default-src 'self' https: data:; "
        "script-src 'self' 'unsafe-inline'; "
        "style-src 'self' 'unsafe-inline' https://cdnjs.cloudflare.com; "
        "img-src 'self' https: data:; "
        "font-src 'self' https://cdnjs.cloudflare.com data:; "
        "connect-src 'self' https:; "
        "frame-ancestors 'self';"
    )
    return response


def _resolve_host(request) -> str:
    """解析真实请求域名。

    优先级：X-Forwarded-Host（反代会保留原始 Host）> request.host。
    取第一个（链式反代时取最左），去掉端口号，转小写。
    """
    fwd = request.headers.get("X-Forwarded-Host", "")
    if fwd:
        host = fwd.split(",")[0].strip().lower()
    else:
        host = request.host.lower()
    return host.split(":")[0]


@web.middleware
async def tenant_middleware(request, handler):
    """多租户中间件：解析租户并挂载对应 DB 连接。

    判定逻辑:
    - /super-admin* 路径 → 任何域名下都走超级管理员上下文
      （未单独绑定主域名时，子租户域名也可访问超管后台）
    - 其余路径：从 Host 解析子域名作为 slug
      · 提取到 slug 且租户存在/启用 → 租户上下文
      · 提取不到 slug（主域名/IP/localhost）→ 超管上下文
    - 单租户模式 → 直接使用 default 租户
    """
    config = request.app["config"]
    tenant_mgr = request.app["tenant_manager"]

    # 单租户兼容模式
    if config.single_tenant:
        request["tenant_slug"] = "default"
        request["tenant_info"] = None
        request["is_super_admin"] = False
        request["db_conn"] = request.app["db_conn"]
        return await handler(request)

    path = request.path

    # 超管路由：任何域名都可访问 /super-admin*
    if path == "/super-admin" or path.startswith("/super-admin/"):
        request["tenant_slug"] = None
        request["tenant_info"] = None
        request["is_super_admin"] = True
        request["db_conn"] = None
        return await handler(request)

    # 其余路由：从 Host 解析租户 slug
    host = _resolve_host(request)

    # 提取子域名第一段
    parts = host.split(".") if host else []
    slug = None
    if len(parts) >= 3:
        slug = parts[0]
        if slug in ("www", ""):
            slug = None

    if slug is None:
        # 主域名 / IP / localhost → 超级管理员上下文
        request["tenant_slug"] = None
        request["tenant_info"] = None
        request["is_super_admin"] = True
        request["db_conn"] = None
        return await handler(request)

    # 多租户模式：查找租户
    tenant_info = await tenant_mgr.get_tenant(slug)
    if tenant_info is None:
        raise web.HTTPNotFound(text=f"租户 '{slug}' 不存在")

    if not tenant_info.get("is_active", 1):
        raise web.HTTPForbidden(text=f"租户 '{slug}' 已禁用")

    request["tenant_slug"] = slug
    request["tenant_info"] = tenant_info
    request["is_super_admin"] = False
    request["db_conn"] = await tenant_mgr.get_tenant_db(slug)
    return await handler(request)


# ---------------------------------------------------------------------------
# 数据库连接管理
# ---------------------------------------------------------------------------
class TenantManager:
    """管理主控库 (tenants.db) 和所有租户数据库连接。"""

    def __init__(self):
        self.master_conn = None
        self._tenant_dbs = {}       # {slug: aiosqlite.Connection}
        self._deleting = set()      # 正在删除的 slug，拒绝新请求

    async def init_master(self):
        """初始化主控数据库，确保 tenants 表存在。"""
        os.makedirs("instance", exist_ok=True)
        self.master_conn = await aiosqlite.connect("instance/tenants.db")
        self.master_conn.row_factory = aiosqlite.Row
        await self.master_conn.execute("""
            CREATE TABLE IF NOT EXISTS tenants (
                slug TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                admin_token TEXT DEFAULT '',
                is_active INTEGER DEFAULT 1,
                created_at TEXT DEFAULT (datetime('now'))
            )
        """)
        await self.master_conn.commit()

    async def get_tenant(self, slug: str) -> dict | None:
        """查询租户信息。"""
        cursor = await self.master_conn.execute(
            "SELECT slug, name, admin_token, is_active, created_at FROM tenants WHERE slug = ?",
            (slug,)
        )
        row = await cursor.fetchone()
        await cursor.close()
        return dict(row) if row else None

    async def list_tenants(self) -> list[dict]:
        """列出所有租户。"""
        cursor = await self.master_conn.execute(
            "SELECT slug, name, admin_token, is_active, created_at FROM tenants ORDER BY created_at DESC"
        )
        rows = await cursor.fetchall()
        await cursor.close()
        return [dict(r) for r in rows]

    async def create_tenant(self, slug: str, name: str, admin_token: str = "") -> dict:
        """创建租户：写入主控库 + 初始化租户数据库。"""
        # 写入主控库
        await self.master_conn.execute(
            "INSERT INTO tenants (slug, name, admin_token, is_active) VALUES (?, ?, ?, 1)",
            (slug, name, admin_token)
        )
        await self.master_conn.commit()

        # 初始化租户数据库
        await self._init_tenant_db(slug, name)

        return await self.get_tenant(slug)

    async def delete_tenant(self, slug: str):
        """删除租户：标记删除中 → 移除主控库记录 → 关闭连接 → 删除文件。

        标记 _deleting 后，新的 get_tenant_db 调用会立即拒绝，
        避免 in-flight 请求在 close/remove 之后仍使用连接。
        """
        self._deleting.add(slug)
        try:
            await self.master_conn.execute("DELETE FROM tenants WHERE slug = ?", (slug,))
            await self.master_conn.commit()

            # 关闭并移除连接
            if slug in self._tenant_dbs:
                await self._tenant_dbs[slug].close()
                del self._tenant_dbs[slug]

            # 删除数据库文件
            db_path = f"instance/{slug}.db"
            if os.path.exists(db_path):
                os.remove(db_path)

            # 删除上传目录（递归）
            uploads_dir = f"static/uploads/{slug}"
            if os.path.isdir(uploads_dir):
                import shutil
                shutil.rmtree(uploads_dir, ignore_errors=True)
        finally:
            self._deleting.discard(slug)

    async def toggle_tenant(self, slug: str, is_active: bool):
        """启用/禁用租户。"""
        await self.master_conn.execute(
            "UPDATE tenants SET is_active = ? WHERE slug = ?",
            (1 if is_active else 0, slug)
        )
        await self.master_conn.commit()

    async def update_tenant_token(self, slug: str, new_token: str):
        """更新租户的 admin_token。"""
        await self.master_conn.execute(
            "UPDATE tenants SET admin_token = ? WHERE slug = ?",
            (new_token, slug)
        )
        await self.master_conn.commit()

    async def rename_tenant(self, old_slug: str, new_slug: str):
        """重命名租户 slug（同步更新 DB 文件、上传目录、site_settings 中的路径）。"""
        if old_slug == new_slug:
            return
        # 检查新 slug 是否已存在
        existing = await self.get_tenant(new_slug)
        if existing:
            raise ValueError(f"租户 '{new_slug}' 已存在")

        # 1. 更新主控库
        await self.master_conn.execute(
            "UPDATE tenants SET slug = ? WHERE slug = ?",
            (new_slug, old_slug)
        )
        await self.master_conn.commit()

        # 2. 关闭旧连接
        if old_slug in self._tenant_dbs:
            await self._tenant_dbs[old_slug].close()
            del self._tenant_dbs[old_slug]

        # 3. 重命名数据库文件
        old_db = f"instance/{old_slug}.db"
        new_db = f"instance/{new_slug}.db"
        if os.path.exists(old_db):
            os.rename(old_db, new_db)

        # 4. 更新 site_settings 中的上传路径（把 /uploads/{old_slug}/ 替换为 /uploads/{new_slug}/）
        conn = await aiosqlite.connect(new_db)
        conn.row_factory = aiosqlite.Row
        cursor = await conn.execute("SELECT key, value FROM site_settings")
        rows = await cursor.fetchall()
        await cursor.close()
        for row in rows:
            key, value = row["key"], row["value"]
            if value and f"/uploads/{old_slug}/" in value:
                new_value = value.replace(f"/uploads/{old_slug}/", f"/uploads/{new_slug}/")
                await conn.execute(
                    "UPDATE site_settings SET value = ? WHERE key = ?",
                    (new_value, key)
                )
        await conn.commit()
        await conn.close()

        # 5. 重命名上传目录
        old_uploads = f"static/uploads/{old_slug}"
        new_uploads = f"static/uploads/{new_slug}"
        if os.path.isdir(old_uploads):
            os.rename(old_uploads, new_uploads)

    async def get_tenant_db(self, slug: str) -> aiosqlite.Connection:
        """获取或创建租户数据库连接（懒加载）。"""
        if slug in self._deleting:
            raise web.HTTPNotFound(text=f"租户 '{slug}' 正在删除")
        if slug not in self._tenant_dbs:
            db_path = f"instance/{slug}.db"
            conn = await aiosqlite.connect(db_path)
            conn.row_factory = aiosqlite.Row
            await self._ensure_tables(conn)
            self._tenant_dbs[slug] = conn
        return self._tenant_dbs[slug]

    async def _init_tenant_db(self, slug: str, name: str):
        """创建新租户的数据库并写入默认设置。"""
        db_path = f"instance/{slug}.db"
        os.makedirs("instance", exist_ok=True)
        conn = await aiosqlite.connect(db_path)
        conn.row_factory = aiosqlite.Row
        await self._ensure_tables(conn)

        # 写入该租户的默认站点设置
        defaults = DEFAULT_SETTINGS.copy()
        defaults["title"] = f"{name}的歌单"
        defaults["singer_name"] = name
        for key, value in defaults.items():
            await conn.execute(
                "INSERT INTO site_settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                (key, value),
            )
        await conn.commit()
        await conn.close()

    async def _ensure_tables(self, conn):
        """创建租户数据库表。"""
        await conn.execute("""
            CREATE TABLE IF NOT EXISTS songs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                artist TEXT NOT NULL,
                language TEXT NOT NULL,
                genre TEXT NOT NULL,
                url TEXT NOT NULL
            )
        """)
        await conn.execute("""
            CREATE TABLE IF NOT EXISTS site_settings (
                key TEXT PRIMARY KEY,
                value TEXT
            )
        """)
        await conn.commit()

    async def preload_tenants(self):
        """预加载所有活跃租户的数据库连接。"""
        tenants = await self.list_tenants()
        for t in tenants:
            if t.get("is_active", 1):
                await self.get_tenant_db(t["slug"])

    async def close_all(self):
        """关闭所有数据库连接。"""
        for conn in self._tenant_dbs.values():
            await conn.close()
        self._tenant_dbs.clear()
        if self.master_conn:
            await self.master_conn.close()


async def create_db_connection(db_path):
    """创建独立的数据库连接（用于单租户模式）。"""
    db_dir = os.path.dirname(db_path)
    if db_dir:
        os.makedirs(db_dir, exist_ok=True)
    conn = await aiosqlite.connect(db_path)
    conn.row_factory = aiosqlite.Row
    await conn.execute("""
        CREATE TABLE IF NOT EXISTS songs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            artist TEXT NOT NULL,
            language TEXT NOT NULL,
            genre TEXT NOT NULL,
            url TEXT NOT NULL
        )
    """)
    await conn.execute("""
        CREATE TABLE IF NOT EXISTS site_settings (
            key TEXT PRIMARY KEY,
            value TEXT
        )
    """)
    await conn.commit()
    return conn


# ---------------------------------------------------------------------------
# 文件上传辅助
# ---------------------------------------------------------------------------
def ensure_upload_dir(slug=None):
    """确保上传目录存在。多租户时按 slug 隔离。"""
    if slug:
        path = f"static/uploads/{slug}"
    else:
        path = "static/uploads"
    os.makedirs(path, exist_ok=True)
    return path


def save_file_field(file_field, prefix, slug=None):
    """Save aiohttp FileField 到上传目录，返回 web 路径。"""
    filename = file_field.filename
    if not filename:
        return None
    ext = os.path.splitext(filename)[1]
    safe_name = f"{prefix}_{int(time.time())}{ext}"
    upload_dir = ensure_upload_dir(slug)
    path_fs = os.path.join(upload_dir, safe_name)
    with open(path_fs, "wb") as f:
        content = file_field.file.read()
        f.write(content)
    if slug:
        return f"/static/uploads/{slug}/{safe_name}"
    return f"/static/uploads/{safe_name}"


# ---------------------------------------------------------------------------
# 字体
# ---------------------------------------------------------------------------
def pick_font(font_path=None, size=38):
    """Pick a font that exists on the system."""
    # 安全：禁止绝对路径与目录穿越，避免探测任意文件是否存在
    if font_path:
        if os.path.isabs(font_path) or ".." in font_path.split("/"):
            font_path = None
    env_font = os.environ.get("QQZHU_FONT_PATH")
    candidates = []
    if font_path:
        candidates.append(font_path)
    if env_font:
        candidates.append(env_font)
    candidates.append(os.path.join("static", "simsun.ttc"))
    candidates.extend([
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansSC-Regular.otf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    ])
    candidates.extend([
        "msyhl.ttc", "SourceHanSansSC-Light.otf", "msyh.ttc",
        "simsun.ttc", "STHeiti Light.ttc",
    ])
    for path in candidates:
        if not path or not os.path.exists(path):
            continue
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            continue
    font = ImageFont.load_default()
    font.size = size
    return font


# ---------------------------------------------------------------------------
# 歌单图片生成
# ---------------------------------------------------------------------------
def wrap_song_names(names, max_chars_per_line=35, max_songs_per_line=6):
    lines = []
    current = []
    current_chars = 0
    for name in names:
        name_len = len(name)
        if current and (current_chars + name_len > max_chars_per_line or len(current) >= max_songs_per_line):
            lines.append("  ".join(current))
            current = []
            current_chars = 0
        current.append(name)
        current_chars += name_len + 2
    if current:
        lines.append("  ".join(current))
    return lines


def combine_background_img(head_img, content_img, end_img, output_height):
    head_h = head_img.height
    content_h = content_img.height
    end_h = end_img.height
    width = head_img.width
    combined = Image.new("RGB", (width, output_height))
    combined.paste(head_img, (0, 0))
    content_area_h = output_height - head_h - end_h
    blocks = max(content_area_h // content_h, 0)
    for i in range(blocks + 1):
        y = head_h + i * content_h
        if y >= output_height - end_h:
            break
        combined.paste(content_img, (0, y))
    combined.paste(end_img, (0, output_height - end_h))
    return combined


def measure_text(draw_obj, text, font):
    """Measure text size; compatible with Pillow versions where textsize is removed."""
    if hasattr(draw_obj, "textbbox"):
        bbox = draw_obj.textbbox((0, 0), text, font=font)
        return bbox[2] - bbox[0], bbox[3] - bbox[1]
    if hasattr(draw_obj, "textsize"):
        return draw_obj.textsize(text, font=font)
    return font.getsize(text)


def generate_playlist_image_from_bg(
    bg_bytes, content_start, end_start, names,
    font_path=None, font_size=38,
    max_chars_per_line=35, max_songs_per_line=6,
    line_height=None, output_dir="static/uploads",
):
    if not names:
        raise ValueError("歌曲列表为空")
    os.makedirs(output_dir, exist_ok=True)
    img = Image.open(io.BytesIO(bg_bytes))
    width, height = img.size
    if not (0 <= content_start < end_start <= height):
        raise ValueError("content_start/end_start 范围非法")

    head = img.crop((0, 0, width, content_start))
    content = img.crop((0, content_start, width, end_start))
    tail = img.crop((0, end_start, width, height))

    bg_color = content.getpixel((min(10, content.width - 1), min(10, content.height - 1)))
    text_color = (30, 30, 30) if sum(bg_color) > 382 else (240, 240, 240)

    lines = wrap_song_names(names, max_chars_per_line=max_chars_per_line, max_songs_per_line=max_songs_per_line)
    resolved_line_height = max(line_height or content.height, 80)
    total_height = head.height + len(lines) * resolved_line_height + tail.height
    canvas = combine_background_img(head, content, tail, total_height)

    font = pick_font(font_path, font_size)
    draw = ImageDraw.Draw(canvas)
    shadow_color = (
        (bg_color[0] - 30, bg_color[1] - 30, bg_color[2] - 30)
        if sum(bg_color) > 382 else (0, 0, 0)
    )

    y = head.height
    margin_left = 50
    for line in lines:
        text_w, text_h = measure_text(draw, line, font)
        text_y = y + (resolved_line_height - text_h) // 2
        draw.text((margin_left + 2, text_y + 2), line, font=font, fill=shadow_color)
        draw.text((margin_left, text_y), line, font=font, fill=text_color)
        y += resolved_line_height

    ts = int(time.time())
    filename = f"playlist_{ts}.png"
    path_fs = os.path.join(output_dir, filename)
    canvas.save(path_fs, quality=95)
    web_path = f"/{output_dir.replace(os.sep, '/')}/{filename}"
    return web_path


def generate_playlist_pages_from_bg(
    bg_bytes, rect, names,
    font_path=None, font_size=38,
    max_chars_per_line=35, max_songs_per_line=6,
    line_height=80, output_dir="static/uploads",
):
    if not names:
        raise ValueError("歌曲列表为空")
    os.makedirs(output_dir, exist_ok=True)
    img = Image.open(io.BytesIO(bg_bytes)).convert("RGB")
    width, height = img.size
    x1, y1, x2, y2 = rect
    if not (0 <= x1 < x2 <= width and 0 <= y1 < y2 <= height):
        raise ValueError("矩形范围非法")

    lines = wrap_song_names(names, max_chars_per_line=max_chars_per_line, max_songs_per_line=max_songs_per_line)
    lines_per_page = max(1, (y2 - y1) // line_height)
    pages = math.ceil(len(lines) / lines_per_page)

    sample_x = min(max(x1 + 5, 0), width - 1)
    sample_y = min(max(y1 + 5, 0), height - 1)
    bg_color = img.getpixel((sample_x, sample_y))
    text_color = (30, 30, 30) if sum(bg_color) > 382 else (240, 240, 240)
    shadow_color = (
        (max(bg_color[0] - 30, 0), max(bg_color[1] - 30, 0), max(bg_color[2] - 30, 0))
        if sum(bg_color) > 382 else (0, 0, 0)
    )
    font = pick_font(font_path, font_size)

    ts = int(time.time())
    saved_files = []
    for page in range(pages):
        start = page * lines_per_page
        end = start + lines_per_page
        page_lines = lines[start:end]
        canvas = img.copy()
        draw = ImageDraw.Draw(canvas)
        y = y1
        for line in page_lines:
            text_w, text_h = measure_text(draw, line, font)
            text_y = y + (line_height - text_h) // 2
            draw.text((x1 + 2, text_y + 2), line, font=font, fill=shadow_color)
            draw.text((x1, text_y), line, font=font, fill=text_color)
            y += line_height
        filename = f"playlist_page_{ts}_{page+1:02d}.png"
        path_fs = os.path.join(output_dir, filename)
        canvas.save(path_fs, quality=95)
        saved_files.append(path_fs)

    zip_name = f"playlist_pages_{ts}.zip"
    zip_path = os.path.join(output_dir, zip_name)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for fpath in saved_files:
            zf.write(fpath, arcname=os.path.basename(fpath))
    web_prefix = f"/{output_dir.replace(os.sep, '/')}/"
    return {
        "zip_path": f"{web_prefix}{zip_name}",
        "files": [f"{web_prefix}{os.path.basename(p)}" for p in saved_files],
    }


# ---------------------------------------------------------------------------
# 辅助
# ---------------------------------------------------------------------------
def wants_json(request):
    accept = request.headers.get("Accept", "")
    return "application/json" in accept or request.headers.get("X-Requested-With") == "XMLHttpRequest"


async def fetch_songs(conn):
    cursor = await conn.execute(
        "SELECT id, name, artist, language, genre, url FROM songs ORDER BY id ASC"
    )
    rows = await cursor.fetchall()
    await cursor.close()
    return [dict(row) for row in rows]


async def fetch_songs_sorted(conn):
    songs = await fetch_songs(conn)

    def sort_key(song):
        name = song["name"]
        language = song["language"]
        language_priority = 0 if language == "中文" else 1
        word_count = len(name)
        if language == "中文":
            name_for_sort = "".join([item[0] for item in pinyin(name, style=Style.NORMAL)])
        else:
            name_for_sort = name.lower()
        return (language_priority, word_count, name_for_sort)

    return sorted(songs, key=sort_key)


async def fetch_unique_languages(conn):
    cursor = await conn.execute("SELECT DISTINCT language FROM songs")
    rows = await cursor.fetchall()
    await cursor.close()
    return [row["language"] for row in rows]


async def fetch_unique_genres(conn):
    cursor = await conn.execute("SELECT DISTINCT genre FROM songs")
    rows = await cursor.fetchall()
    await cursor.close()
    return [row["genre"] for row in rows]


async def get_settings(conn):
    cursor = await conn.execute("SELECT key, value FROM site_settings")
    rows = await cursor.fetchall()
    await cursor.close()
    stored = {row["key"]: row["value"] for row in rows}
    merged = DEFAULT_SETTINGS.copy()
    merged.update({k: v for k, v in stored.items() if v is not None})
    return merged


async def update_settings(conn, settings: dict):
    for key, value in settings.items():
        await conn.execute(
            "INSERT INTO site_settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value",
            (key, value),
        )
    await conn.commit()


async def add_song(conn, name, artist, language, genre, url):
    cursor = await conn.execute(
        "INSERT INTO songs (name, artist, language, genre, url) VALUES (?, ?, ?, ?, ?)",
        (name, artist, language, genre, url),
    )
    song_id = cursor.lastrowid
    await conn.commit()
    await cursor.close()
    return song_id


async def update_song(conn, song_id, name, artist, language, genre, url):
    await conn.execute(
        "UPDATE songs SET name = ?, artist = ?, language = ?, genre = ?, url = ? WHERE id = ?",
        (name, artist, language, genre, url, song_id),
    )
    await conn.commit()


async def delete_song(conn, song_id):
    await conn.execute("DELETE FROM songs WHERE id = ?", (song_id,))
    await conn.commit()


async def backup_songs(conn, dest_path):
    songs = await fetch_songs(conn)
    os.makedirs(os.path.dirname(dest_path), exist_ok=True)
    with open(dest_path, "w", encoding="utf-8") as f:
        json.dump(songs, f, ensure_ascii=False, indent=2)
    return dest_path


def parse_backup_json(text: str):
    data = json.loads(text)
    if not isinstance(data, list):
        raise ValueError("备份格式错误：根节点应为列表")
    for item in data:
        if not isinstance(item, dict):
            raise ValueError("备份格式错误：列表元素应为对象")
        for key in ("name", "artist", "language", "genre", "url"):
            if key not in item:
                raise ValueError(f"备份缺少必要字段: {key}")
    return data


async def restore_songs_from_data(conn, songs):
    await conn.execute("DELETE FROM songs")
    for s in songs:
        await add_song(
            conn,
            s.get("name", ""), s.get("artist", ""),
            s.get("language", ""), s.get("genre", ""), s.get("url", "-"),
        )
    await conn.commit()
    return len(songs)


def parse_backup_xlsx(data: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError("未安装 openpyxl，无法解析 xlsx，请先安装 openpyxl")

    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    required = ["歌名", "歌手", "语言", "风格", "url"]
    if header[:5] != required:
        raise ValueError(f"xlsx 表头需为：{' | '.join(required)}")
    songs = []
    for row in rows[1:]:
        if row is None:
            continue
        vals = list(row) + [""] * (5 - len(row))
        name, artist, language, genre, url = ["" if v is None else str(v) for v in vals[:5]]
        if not name and not artist:
            continue
        songs.append({"name": name, "artist": artist, "language": language, "genre": genre, "url": url or "-"})
    return songs


# ---------------------------------------------------------------------------
# 鉴权
# ---------------------------------------------------------------------------
def require_admin(request, form=None):
    """校验管理员 token。多租户模式下检查租户的 admin_token；单租户模式下检查全局 token。"""
    config = request.app["config"]
    tenant_info = request.get("tenant_info")

    if tenant_info is not None:
        # 多租户模式：使用租户自己的 admin_token
        token_cfg = tenant_info.get("admin_token", "")
    else:
        # 单租户模式
        token_cfg = config.admin_token()

    if not token_cfg:
        next_url = quote(str(request.rel_url))
        raise web.HTTPFound(location=f"/admin/setup?next={next_url}")

    provided = (
        request.headers.get("X-Admin-Token")
        or request.query.get("token")
        or (form.get("token") if form else None)
        or request.cookies.get("admin_token")
    )
    if token_cfg and secrets.compare_digest(str(provided or ""), str(token_cfg)):
        return provided
    next_url = quote(str(request.rel_url))
    raise web.HTTPFound(location=f"/admin/login?next={next_url}")


def require_super_admin(request, form=None):
    """校验超级管理员 token。"""
    config = request.app["config"]
    token_cfg = config.super_admin_token()
    if not token_cfg:
        # 未设置超级管理员 token，允许无认证访问（首次设置）
        return True

    provided = (
        request.headers.get("X-Super-Admin-Token")
        or request.query.get("token")
        or (form.get("token") if form else None)
        or request.cookies.get("super_admin_token")
    )
    if token_cfg and secrets.compare_digest(str(provided or ""), str(token_cfg)):
        return True
    next_url = quote(str(request.rel_url))
    raise web.HTTPFound(location=f"/super-admin/login?next={next_url}")


# ---------------------------------------------------------------------------
# 租户端 - 前台
# ---------------------------------------------------------------------------
async def index(request):
    conn = request["db_conn"]
    # 主域名 / localhost / 纯 IP → 没有租户上下文，跳转到超级管理员后台
    if conn is None:
        raise web.HTTPFound(location="/super-admin")
    songs = await fetch_songs_sorted(conn)
    languages = await fetch_unique_languages(conn)
    genres = await fetch_unique_genres(conn)
    settings = await get_settings(conn)
    return aiohttp_jinja2.render_template(
        "index.html", request,
        {"songs": songs, "languages": languages, "genres": genres, "settings": settings},
    )


async def _check_song_playable(session, song_id, headers):
    """检测指定网易云歌曲 ID 是否存在有效音频流（过滤 404 及 HTML 错误页）。"""
    outer_url = f"https://music.163.com/song/media/outer/url?id={song_id}.mp3"
    try:
        async with session.get(outer_url, headers=headers, allow_redirects=True, timeout=aiohttp.ClientTimeout(total=4)) as resp:
            final_url = str(resp.url)
            content_type = resp.headers.get("Content-Type", "")
            if "404" not in final_url and "text/html" not in content_type:
                return True, final_url
    except Exception:
        pass
    return False, outer_url


async def netease_preview_handler(request):
    """根据歌名和歌手匹配网易云音乐，自动检测音源有效性并返回试听链接与元数据。"""
    name = request.query.get("name", "").strip()
    artist = request.query.get("artist", "").strip()
    if not name:
        return web.json_response({"ok": False, "message": "歌名不能为空"})

    url = "https://music.163.com/api/cloudsearch/pc"
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Referer": "https://music.163.com/",
        "Cookie": "appver=2.0.2; NMTID=00h; os=pc;",
    }

    def parse_song_meta(song):
        song_name = song.get("name", name)
        ar_objs = song.get("ar", song.get("artists", [])) or []
        artists_list = [ar.get("name") for ar in ar_objs if isinstance(ar, dict) and ar.get("name")]
        artist_name = " / ".join(artists_list) if artists_list else artist
        album = song.get("al", song.get("album", {})) or {}
        pic_url = album.get("picUrl", "") if isinstance(album, dict) else ""
        return song_name, artist_name, pic_url

    try:
        timeout = aiohttp.ClientTimeout(total=8)
        async with aiohttp.ClientSession(timeout=timeout) as session:
            # 1. 优先尝试 歌名 + 歌手 检索
            keywords = f"{name} {artist}".strip()
            form_data = {"s": keywords, "type": "1", "offset": "0", "limit": "10"}
            async with session.post(url, data=form_data, headers=headers) as resp:
                if resp.status != 200:
                    return web.json_response({"ok": False, "message": f"网易云 API 请求失败 ({resp.status})"})
                data = await resp.json(content_type=None)
                songs = (data.get("result", {}) or {}).get("songs", []) or []

            # 2. 如果没搜索到结果，降级为仅按 歌名 检索
            if not songs and artist:
                form_data["s"] = name
                async with session.post(url, data=form_data, headers=headers) as resp:
                    if resp.status == 200:
                        data = await resp.json(content_type=None)
                        songs = (data.get("result", {}) or {}).get("songs", []) or []

            if not songs:
                return web.json_response({"ok": False, "message": "网易云未找到匹配歌曲"})

            # 3. 按歌名完全匹配排序（同名优先）
            name_clean = name.strip().lower()
            exact_matches = [s for s in songs if s.get("name", "").strip().lower() == name_clean]
            other_matches = [s for s in songs if s.get("name", "").strip().lower() != name_clean]
            candidate_list = exact_matches + other_matches

            # 4. 遍历候选列表，检测音频外链可用性（自动过滤无版权 404）
            for song in candidate_list:
                song_id = song.get("id")
                playable, audio_url = await _check_song_playable(session, song_id, headers)
                if playable:
                    song_name, artist_name, pic_url = parse_song_meta(song)
                    return web.json_response({
                        "ok": True,
                        "song_id": song_id,
                        "audio_url": audio_url,
                        "name": song_name,
                        "artist": artist_name,
                        "picUrl": pic_url
                    })

            # 5. 如果 歌名+歌手 搜出的候选均无可用音源，再用 纯歌名 广域搜一次候选
            if artist:
                form_data["s"] = name
                async with session.post(url, data=form_data, headers=headers) as resp:
                    if resp.status == 200:
                        data = await resp.json(content_type=None)
                        fallback_songs = (data.get("result", {}) or {}).get("songs", []) or []
                
                fallback_exact = [s for s in fallback_songs if s.get("name", "").strip().lower() == name_clean]
                fallback_list = fallback_exact + [s for s in fallback_songs if s not in fallback_exact]
                for song in fallback_list:
                    song_id = song.get("id")
                    playable, audio_url = await _check_song_playable(session, song_id, headers)
                    if playable:
                        song_name, artist_name, pic_url = parse_song_meta(song)
                        return web.json_response({
                            "ok": True,
                            "song_id": song_id,
                            "audio_url": audio_url,
                            "name": song_name,
                            "artist": artist_name,
                            "picUrl": pic_url
                        })

            return web.json_response({"ok": False, "message": "该歌曲在网易云暂无可用试听资源（受版权保护或已下架）"})
    except Exception as exc:
        return web.json_response({"ok": False, "message": f"获取试听失败: {exc}"})


async def proxy_image(request):
    image_url = request.query.get("url")
    if not image_url:
        raise web.HTTPBadRequest(text="Missing 'url' parameter")

    # SSRF 防护：仅允许 http/https + 非内网 IP + bilibili 图片域名
    parsed = urlparse(image_url)
    if parsed.scheme not in ("http", "https"):
        raise web.HTTPBadRequest(text="仅支持 http/https 协议")
    hostname = parsed.hostname
    if not hostname:
        raise web.HTTPBadRequest(text="无效的 URL")
    try:
        ip = ipaddress.ip_address(hostname)
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
            raise web.HTTPForbidden(text="禁止访问内网/保留地址")
    except ValueError:
        pass  # 域名，继续做域名白名单校验
    allowed_domains = ("hdslb.com", "bilibili.com", "bilivideo.com", "bilivideo.cn")
    if not any(hostname == d or hostname.endswith("." + d) for d in allowed_domains):
        raise web.HTTPForbidden(text="仅允许 bilibili 相关图片域名")

    headers = {"Referer": "https://www.bilibili.com"}
    async with aiohttp.ClientSession() as session:
        async with session.get(image_url, headers=headers) as response:
            if response.status == 200:
                return web.Response(
                    body=await response.read(), content_type=response.headers["Content-Type"]
                )
            raise web.HTTPNotFound(text="Image not found")


# ---------------------------------------------------------------------------
# 租户端 - 后台认证
# ---------------------------------------------------------------------------
async def admin_login_get(request):
    config = request.app["config"]
    tenant_info = request.get("tenant_info")
    if tenant_info is not None:
        token_cfg = tenant_info.get("admin_token", "")
    else:
        token_cfg = config.admin_token()

    if not token_cfg:
        next_url = quote(str(request.rel_url))
        raise web.HTTPFound(location=f"/admin/setup?next={next_url}")

    # 获取租户背景图设置
    background_url = "/static/background.jpg"
    if request.get("db_conn"):
        try:
            settings = await get_settings(request["db_conn"])
            background_url = settings.get("background_url", background_url)
        except Exception:
            pass

    return aiohttp_jinja2.render_template(
        "admin_login.html", request,
        {
            "next": request.query.get("next", "/admin"),
            "message": request.query.get("message", ""),
            "csrf_token": request.get("csrf_token", ""),
            "background_url": background_url,
        },
    )


async def admin_login_post(request):
    ip = (request.headers.get("X-Forwarded-For", "").split(",")[0].strip() or request.remote or "unknown")
    remaining, wait = _login_attempts_info(ip)

    # 获取租户背景图设置
    background_url = "/static/background.jpg"
    if request.get("db_conn"):
        try:
            settings = await get_settings(request["db_conn"])
            background_url = settings.get("background_url", background_url)
        except Exception:
            pass

    if remaining == 0:
        return aiohttp_jinja2.render_template(
            "admin_login.html", request,
            {
                "next": request.query.get("next", "/admin"),
                "message": f"尝试过多，请 {wait} 秒后再试",
                "csrf_token": request.get("csrf_token", ""),
                "background_url": background_url,
            },
        )

    config = request.app["config"]
    tenant_info = request.get("tenant_info")
    if tenant_info is not None:
        token_cfg = tenant_info.get("admin_token", "")
    else:
        token_cfg = config.admin_token()

    if not token_cfg:
        next_url = quote(str(request.rel_url))
        raise web.HTTPFound(location=f"/admin/setup?next={next_url}")

    form = await request.post()
    provided = form.get("token", "")
    next_url = form.get("next") or "/admin"

    if token_cfg and secrets.compare_digest(str(provided or ""), str(token_cfg)):
        resp = web.HTTPFound(location=next_url)
        resp.set_cookie("admin_token", provided, httponly=True, samesite="Lax")
        _reset_login_failure(ip)
        return resp

    _add_login_failure(ip)
    remaining_after, wait_after = _login_attempts_info(ip)
    message_text = "Token 错误或未设置"
    if remaining_after > 0:
        message_text += f"；还可再试 {remaining_after} 次"
    else:
        message_text += f"；请 {wait_after} 秒后再试"
    return aiohttp_jinja2.render_template(
        "admin_login.html", request,
        {"next": next_url, "message": message_text, "csrf_token": request.get("csrf_token", ""), "background_url": background_url},
    )


async def admin_logout(request):
    resp = web.HTTPFound(location="/admin/login")
    resp.del_cookie("admin_token")
    return resp


async def admin_setup_get(request):
    """首次设置后台 token（仅当未设置时访问）。"""
    config = request.app["config"]
    tenant_info = request.get("tenant_info")
    if tenant_info is not None:
        token_cfg = tenant_info.get("admin_token", "")
    else:
        token_cfg = config.admin_token()

    if token_cfg:
        raise web.HTTPFound(location="/admin/login")
    return aiohttp_jinja2.render_template(
        "admin_setup.html", request,
        {
            "next": request.query.get("next", "/admin"),
            "message": request.query.get("message", ""),
            "csrf_token": request.get("csrf_token", ""),
        },
    )


async def admin_setup_post(request):
    config = request.app["config"]
    tenant_info = request.get("tenant_info")

    if tenant_info is not None:
        token_cfg = tenant_info.get("admin_token", "")
    else:
        token_cfg = config.admin_token()

    if token_cfg:
        raise web.HTTPFound(location="/admin/login")

    form = await request.post()
    new_token = form.get("new_token", "").strip()
    confirm_token = form.get("confirm_token", "").strip()
    next_url = form.get("next") or "/admin"

    if not new_token:
        return aiohttp_jinja2.render_template(
            "admin_setup.html", request,
            {"next": next_url, "message": "Token 不能为空", "csrf_token": request.get("csrf_token", "")},
        )
    if new_token != confirm_token:
        return aiohttp_jinja2.render_template(
            "admin_setup.html", request,
            {"next": next_url, "message": "两次输入的 Token 不一致", "csrf_token": request.get("csrf_token", "")},
        )

    if tenant_info is not None:
        # 多租户模式：更新 tenants 表
        slug = request["tenant_slug"]
        await request.app["tenant_manager"].update_tenant_token(slug, new_token)
        tenant_info["admin_token"] = new_token
    else:
        # 单租户模式：写入 .env
        update_env_var("QQZHU_ADMIN_TOKEN", new_token)
        config.env_admin_token = new_token

    resp = web.HTTPFound(location=next_url)
    resp.set_cookie("admin_token", new_token, httponly=True, samesite="Lax")
    return resp


# ---------------------------------------------------------------------------
# 租户端 - 后台管理
# ---------------------------------------------------------------------------
async def admin_page(request):
    _ = require_admin(request)
    conn = request["db_conn"]
    songs = await fetch_songs(conn)
    settings = await get_settings(conn)
    return aiohttp_jinja2.render_template(
        "admin.html", request,
        {
            "songs": songs,
            "settings": settings,
            "message": request.query.get("message", ""),
            "token": request.query.get("token", ""),
            "env_admin_token": bool(request.app["config"].env_admin_token),
            "csrf_token": request.get("csrf_token", ""),
        },
    )


async def admin_action(request):
    # 先鉴权，再做任何业务（cookie/header 已足够校验）
    require_admin(request)
    conn = request["db_conn"]
    slug = request.get("tenant_slug")
    form = await request.post()
    action = form.get("action")
    message = ""

    # 按租户隔离的上传目录
    upload_dir = ensure_upload_dir(slug)

    try:
        if action == "song_new":
            name = form.get("song_name", "").strip()
            artist = form.get("artist", "").strip()
            language = form.get("language", "").strip()
            genre = form.get("genre", "").strip()
            url = form.get("url", "").strip() or "-"
            new_id = await add_song(conn, name, artist, language, genre, url)
            message = "歌曲已添加"
            if wants_json(request):
                return web.json_response({"ok": True, "action": "song_new", "song": {
                    "id": new_id, "name": name, "artist": artist, "language": language, "genre": genre, "url": url
                }})

        elif action == "song_update":
            song_id = int(form.get("song_id", 0))
            name = form.get("song_name", "").strip()
            artist = form.get("artist", "").strip()
            language = form.get("language", "").strip()
            genre = form.get("genre", "").strip()
            url = form.get("url", "").strip() or "-"
            await update_song(conn, song_id, name, artist, language, genre, url)
            message = "歌曲已更新"
            if wants_json(request):
                return web.json_response({"ok": True, "action": "song_update", "song_id": song_id})

        elif action == "song_delete":
            song_id = int(form.get("song_id", 0))
            await delete_song(conn, song_id)
            message = "歌曲已删除"
            if wants_json(request):
                return web.json_response({"ok": True, "action": "song_delete", "song_id": song_id})

        elif action == "backup":
            backup_dir = f"backup/{slug}" if slug else "backup"
            os.makedirs(backup_dir, exist_ok=True)
            dest = os.path.join(backup_dir, "songs_backup.json")
            saved = await backup_songs(conn, dest)
            message = f"备份已生成: {saved}"

        elif action == "restore_backup":
            file_field = form.get("backup_file")
            try:
                if file_field and hasattr(file_field, "file") and file_field.filename:
                    filename = file_field.filename.lower()
                    if filename.endswith(".json"):
                        content = file_field.file.read().decode("utf-8")
                        songs = parse_backup_json(content)
                    elif filename.endswith(".xlsx"):
                        content = file_field.file.read()
                        songs = parse_backup_xlsx(content)
                    else:
                        raise ValueError("仅支持 .json 或 .xlsx 备份文件")
                    count = await restore_songs_from_data(conn, songs)
                    message = f"已从上传的备份恢复 {count} 首歌曲"
                    if wants_json(request):
                        return web.json_response({"ok": True, "action": "restore_backup", "count": count})
                else:
                    backup_dir = f"backup/{slug}" if slug else "backup"
                    src = os.path.join(backup_dir, "songs_backup.json")
                    if not os.path.exists(src):
                        raise FileNotFoundError(f"{src} 不存在")
                    with open(src, "r", encoding="utf-8") as f:
                        content = f.read()
                    songs = parse_backup_json(content)
                    count = await restore_songs_from_data(conn, songs)
                    message = f"已从本地备份恢复 {count} 首歌曲"
                    if wants_json(request):
                        return web.json_response({"ok": True, "action": "restore_backup", "count": count})
            except Exception as exc:
                message = f"恢复失败: {exc}"

        elif action == "settings":
            current_settings = await get_settings(conn)
            bg_file = form.get("background_file")
            singer_file = form.get("singer_file")
            new_settings = {
                "title": form.get("title", "").strip() or current_settings.get("title", DEFAULT_SETTINGS["title"]),
                "live_url": form.get("live_url", "").strip() or current_settings.get("live_url", DEFAULT_SETTINGS["live_url"]),
                "singer_name": form.get("singer_name", "").strip() or current_settings.get("singer_name", DEFAULT_SETTINGS["singer_name"]),
                "singer_intro": form.get("singer_intro", "").strip() or current_settings.get("singer_intro", DEFAULT_SETTINGS["singer_intro"]),
                "background_url": current_settings.get("background_url", DEFAULT_SETTINGS["background_url"]),
                "singer_url": current_settings.get("singer_url", DEFAULT_SETTINGS["singer_url"]),
            }
            if hasattr(bg_file, "file") and bg_file.filename:
                saved = save_file_field(bg_file, "bg", slug)
                if saved:
                    new_settings["background_url"] = saved
            if hasattr(singer_file, "file") and singer_file.filename:
                saved = save_file_field(singer_file, "singer", slug)
                if saved:
                    new_settings["singer_url"] = saved
            await update_settings(conn, new_settings)
            message = "站点信息已更新"

        elif action == "update_admin_token":
            new_token = form.get("new_token", "").strip()
            confirm_token = form.get("confirm_token", "").strip()
            if not new_token:
                raise ValueError("新 token 不能为空")
            if new_token != confirm_token:
                raise ValueError("两次输入的 token 不一致")

            tenant_info = request.get("tenant_info")
            if tenant_info is not None:
                await request.app["tenant_manager"].update_tenant_token(slug, new_token)
                tenant_info["admin_token"] = new_token
            else:
                update_env_var("QQZHU_ADMIN_TOKEN", new_token)
                request.app["config"].env_admin_token = new_token

            message = "admin_token 已更新，请重新登录"
            if wants_json(request):
                return web.json_response({"ok": True, "action": "update_admin_token", "message": message})

        elif action == "generate_playlist_image":
            bg_field = form.get("bg_image")
            if not (bg_field and hasattr(bg_field, "file") and bg_field.filename):
                raise ValueError("请上传背景图")
            try:
                content_start = int(form.get("content_start", 0))
                end_start = int(form.get("end_start", 0))
            except Exception:
                raise ValueError("content_start/end_start 需要是整数")
            songs_sorted = await fetch_songs_sorted(conn)
            names = [s["name"] for s in songs_sorted]
            font_path = form.get("font_path", "").strip() or None
            font_size = int(form.get("font_size", 38) or 38)
            line_height = form.get("line_height")
            line_height_val = int(line_height) if line_height else None
            max_chars_per_line = int(form.get("max_chars_per_line", 35) or 35)
            max_songs_per_line = int(form.get("max_songs_per_line", 6) or 6)
            generated_path = generate_playlist_image_from_bg(
                bg_field.file.read(), content_start, end_start, names,
                font_path=font_path, font_size=font_size,
                max_chars_per_line=max_chars_per_line, max_songs_per_line=max_songs_per_line,
                line_height=line_height_val, output_dir=upload_dir,
            )
            if wants_json(request):
                return web.json_response({"ok": True, "action": "generate_playlist_image", "path": generated_path})
            fs_path = generated_path.lstrip("/\\")
            fs_path = fs_path if os.path.isabs(fs_path) else os.path.join(os.getcwd(), fs_path)
            if not os.path.exists(fs_path):
                raise web.HTTPNotFound(text="生成的长图不存在")
            filename = os.path.basename(fs_path)
            headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
            return web.FileResponse(path=fs_path, headers=headers)

        elif action == "generate_playlist_pages":
            bg_field = form.get("bg_image_small")
            if not (bg_field and hasattr(bg_field, "file") and bg_field.filename):
                raise ValueError("请上传背景图")
            try:
                x1 = int(form.get("rect_x1", 0))
                y1 = int(form.get("rect_y1", 0))
                x2 = int(form.get("rect_x2", 0))
                y2 = int(form.get("rect_y2", 0))
            except Exception:
                raise ValueError("矩形坐标需要是整数")
            font_size = int(form.get("font_size", 38) or 38)
            line_height = int(form.get("line_height", 80) or 80)
            max_chars_per_line = int(form.get("max_chars_per_line", 35) or 35)
            max_songs_per_line = int(form.get("max_songs_per_line", 6) or 6)
            songs_sorted = await fetch_songs_sorted(conn)
            names = [s["name"] for s in songs_sorted]
            font_path = form.get("font_path", "").strip() or None
            result = generate_playlist_pages_from_bg(
                bg_field.file.read(), (x1, y1, x2, y2), names,
                font_path=font_path, font_size=font_size,
                max_chars_per_line=max_chars_per_line, max_songs_per_line=max_songs_per_line,
                line_height=line_height, output_dir=upload_dir,
            )
            zip_path = result["zip_path"]
            if wants_json(request):
                return web.json_response({"ok": True, "action": "generate_playlist_pages", "zip_path": zip_path, "files": result["files"]})
            fs_path = zip_path.lstrip("/\\")
            fs_path = fs_path if os.path.isabs(fs_path) else os.path.join(os.getcwd(), fs_path)
            if not os.path.exists(fs_path):
                raise web.HTTPNotFound(text="生成的文件不存在")
            filename = os.path.basename(fs_path)
            headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
            return web.FileResponse(path=fs_path, headers=headers)

        else:
            message = "未知操作"
    except Exception as exc:
        message = f"操作失败: {exc}"
        if wants_json(request):
            return web.json_response({"ok": False, "message": message})

    if wants_json(request):
        return web.json_response({"ok": True, "message": message})
    params = [f"message={quote(message)}"]
    return web.HTTPFound(location="/admin?" + "&".join(params) + "#songs")


async def admin_download_backup(request):
    _ = require_admin(request)
    slug = request.get("tenant_slug")
    conn = request["db_conn"]
    backup_dir = f"backup/{slug}" if slug else "backup"
    os.makedirs(backup_dir, exist_ok=True)
    dest = os.path.join(backup_dir, "songs_backup.json")
    saved = await backup_songs(conn, dest)
    headers = {"Content-Disposition": 'attachment; filename="songs_backup.json"'}
    return web.FileResponse(path=saved, headers=headers)


# ---------------------------------------------------------------------------
# 超级管理员
# ---------------------------------------------------------------------------
async def super_admin_login_get(request):
    """超级管理员登录页。"""
    config = request.app["config"]
    if not config.super_admin_token():
        # 未设置超级管理员 token，直接跳转管理页（首次）
        raise web.HTTPFound(location="/super-admin")
    return aiohttp_jinja2.render_template(
        "admin_login.html", request,
        {
            "next": request.query.get("next", "/super-admin"),
            "message": request.query.get("message", ""),
            "csrf_token": request.get("csrf_token", ""),
            "is_super_admin": True,
            "form_action": "/super-admin/login",
            "background_url": "/static/background.jpg",
        },
    )


async def super_admin_login_post(request):
    """超级管理员登录处理。"""
    ip = (request.headers.get("X-Forwarded-For", "").split(",")[0].strip() or request.remote or "unknown")
    remaining, wait = _login_attempts_info(ip)
    if remaining == 0:
        return aiohttp_jinja2.render_template(
            "admin_login.html", request,
            {
                "next": request.query.get("next", "/super-admin"),
                "message": f"尝试过多，请 {wait} 秒后再试",
                "csrf_token": request.get("csrf_token", ""),
                "is_super_admin": True,
                "form_action": "/super-admin/login",
                "background_url": "/static/background.jpg",
            },
        )

    config = request.app["config"]
    token_cfg = config.super_admin_token()
    if not token_cfg:
        raise web.HTTPFound(location="/super-admin")

    form = await request.post()
    provided = form.get("token", "")
    next_url = form.get("next") or "/super-admin"

    if token_cfg and secrets.compare_digest(str(provided or ""), str(token_cfg)):
        resp = web.HTTPFound(location=next_url)
        resp.set_cookie("super_admin_token", provided, httponly=True, samesite="Lax")
        _reset_login_failure(ip)
        return resp

    _add_login_failure(ip)
    remaining_after, wait_after = _login_attempts_info(ip)
    message_text = "Token 错误"
    if remaining_after > 0:
        message_text += f"；还可再试 {remaining_after} 次"
    else:
        message_text += f"；请 {wait_after} 秒后再试"
    return aiohttp_jinja2.render_template(
        "admin_login.html", request,
        {"next": next_url, "message": message_text, "csrf_token": request.get("csrf_token", ""),
         "is_super_admin": True, "form_action": "/super-admin/login", "background_url": "/static/background.jpg"},
    )


async def super_admin_logout(request):
    resp = web.HTTPFound(location="/super-admin/login")
    resp.del_cookie("super_admin_token")
    return resp


async def super_admin_page(request):
    """超级管理员管理页 - 列出所有租户。"""
    _ = require_super_admin(request)
    tenant_mgr = request.app["tenant_manager"]
    tenants = await tenant_mgr.list_tenants()
    return aiohttp_jinja2.render_template(
        "super_admin.html", request,
        {
            "tenants": tenants,
            "message": request.query.get("message", ""),
            "csrf_token": request.get("csrf_token", ""),
        },
    )


async def super_admin_action(request):
    """超级管理员操作：创建/删除/启禁租户。"""
    require_super_admin(request)
    tenant_mgr = request.app["tenant_manager"]
    form = await request.post()
    action = form.get("action")
    message = ""

    try:
        if action == "create_tenant":
            slug = form.get("slug", "").strip().lower()
            name = form.get("name", "").strip()
            if not slug or not name:
                raise ValueError("租户标识和名称不能为空")
            if not slug.isalnum():
                raise ValueError("租户标识只能包含字母和数字")
            existing = await tenant_mgr.get_tenant(slug)
            if existing:
                raise ValueError(f"租户 '{slug}' 已存在")
            await tenant_mgr.create_tenant(slug, name)
            message = f"租户 '{name}' ({slug}) 创建成功"

        elif action == "delete_tenant":
            slug = form.get("slug", "").strip()
            if not slug:
                raise ValueError("请指定要删除的租户")
            await tenant_mgr.delete_tenant(slug)
            message = f"租户 '{slug}' 已删除"

        elif action == "toggle_tenant":
            slug = form.get("slug", "").strip()
            is_active = form.get("is_active", "1") == "1"
            await tenant_mgr.toggle_tenant(slug, is_active)
            status = "启用" if is_active else "禁用"
            message = f"租户 '{slug}' 已{status}"

        elif action == "rename_tenant":
            old_slug = form.get("old_slug", "").strip()
            new_slug = form.get("new_slug", "").strip()
            if not old_slug or not new_slug:
                raise ValueError("新旧标识不能为空")
            if not new_slug.isalnum():
                raise ValueError("新标识只能包含字母和数字")
            await tenant_mgr.rename_tenant(old_slug, new_slug)
            message = f"租户 '{old_slug}' 已重命名为 '{new_slug}'"

        elif action == "setup_super_token":
            new_token = form.get("new_token", "").strip()
            confirm_token = form.get("confirm_token", "").strip()
            if not new_token:
                raise ValueError("Token 不能为空")
            if new_token != confirm_token:
                raise ValueError("两次输入的 Token 不一致")
            update_env_var("QQZHU_SUPER_ADMIN_TOKEN", new_token)
            request.app["config"].env_super_admin_token = new_token
            message = "超级管理员 Token 已设置"

        else:
            message = "未知操作"
    except Exception as exc:
        message = f"操作失败: {exc}"

    params = [f"message={quote(message)}"]
    return web.HTTPFound(location="/super-admin?" + "&".join(params))


# ---------------------------------------------------------------------------
# 初始化应用
# ---------------------------------------------------------------------------
async def init_app():
    app = web.Application(
        client_max_size=10 * 1024 * 1024,
        middlewares=[tenant_middleware, csrf_middleware, security_headers_middleware],
    )
    aiohttp_jinja2.setup(app, loader=jinja2.FileSystemLoader("templates"))

    config = Config("config.ini")
    app["config"] = config

    # 初始化租户管理器
    tenant_mgr = TenantManager()
    await tenant_mgr.init_master()
    app["tenant_manager"] = tenant_mgr

    if config.single_tenant:
        # 单租户兼容模式
        db_conn = await create_db_connection(config.db_path())
        app["db_conn"] = db_conn
        async def close_db(app):
            await app["db_conn"].close()
        app.on_cleanup.append(close_db)
    else:
        # 多租户模式：预加载所有活跃租户
        await tenant_mgr.preload_tenants()
        # 确保 default 租户存在（兼容旧数据）
        default = await tenant_mgr.get_tenant("default")
        if not default:
            await tenant_mgr.create_tenant("default", "默认主播")

        async def close_all_tenant_dbs(app):
            await app["tenant_manager"].close_all()
        app.on_cleanup.append(close_all_tenant_dbs)

    # 静态文件 & 公共路由
    app.router.add_static("/static/", path="static", name="static")
    app.router.add_get("/proxy-image", proxy_image)
    app.router.add_get("/api/netease-preview", netease_preview_handler)
    app.router.add_get("/healthz", lambda request: web.Response(text="ok"))

    # 租户前台
    app.router.add_get("/", index)

    # 租户后台
    app.router.add_get("/admin/login", admin_login_get)
    app.router.add_post("/admin/login", admin_login_post)
    app.router.add_get("/admin/logout", admin_logout)
    app.router.add_get("/admin/setup", admin_setup_get)
    app.router.add_post("/admin/setup", admin_setup_post)
    app.router.add_get("/admin", admin_page)
    app.router.add_post("/admin/action", admin_action)
    app.router.add_get("/admin/download-backup", admin_download_backup)

    # 超级管理员
    app.router.add_get("/super-admin/login", super_admin_login_get)
    app.router.add_post("/super-admin/login", super_admin_login_post)
    app.router.add_get("/super-admin/logout", super_admin_logout)
    app.router.add_get("/super-admin", super_admin_page)
    app.router.add_post("/super-admin/action", super_admin_action)

    return app


if __name__ == "__main__":
    port = Config("config.ini").server_port()
    web.run_app(init_app(), port=port)
